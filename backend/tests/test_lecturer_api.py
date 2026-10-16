"""Phase-5 gate: the lecturer's own flow over a scored batch, and the walls around it.

The pipeline is stubbed just as in phase 4 (`seeded`); what is exercised here is this phase's API.
Two families of assertions matter most: an override is a second row that leaves the system score
alone, because the agreement study depends on holding both (FR-11, PRD §12); and another lecturer's
rows are never reachable, not even by id.
"""

import csv
import io
import uuid

import openpyxl
from conftest import INDICATORS, LEVEL_NAMES, LEVELS, QUOTE_COGNITIVE

from app import db
from app.ingest import identity_store
from app.jobs.queue import claim_jobs
from app.routers.export import COLUMNS, RESEARCH_COLUMNS

OTHER_LECTURER = {
    "email": "lain@example.ac.id",
    "name": "Dosen Lain",
    "password": "kata-sandi-panjang-2026",
}


def _run_batch(seeded, mode: str = "multi_agent") -> str:
    """Evaluate the seeded batch once and return the run id."""
    client, csrf = seeded["client"], seeded["csrf"]
    created = client.post(
        f"/assignments/{seeded['assignment']}/runs?mode={mode}", headers={"X-CSRF-Token": csrf}
    )
    assert created.status_code == 201, created.text
    run_id = created.json()["id"]
    processed = client.post(f"/runs/{run_id}/process?limit=5", headers={"X-CSRF-Token": csrf})
    assert processed.status_code == 200, processed.text
    assert processed.json()["processed"] == 1, processed.text
    return run_id


def _csrf(client) -> dict:
    return {"X-CSRF-Token": client.cookies.get("ta_csrf")}


def _score_id(user, run_id: str, dimension: str = "cognitive") -> uuid.UUID:
    with db.tenant_tx(user) as conn:
        return conn.execute(
            "select id from system_scores where run_id = %s and dimension = %s",
            (run_id, dimension),
        ).fetchone()[0]


def _rows_for(conn, submission_id: uuid.UUID) -> dict[str, int]:
    """How many rows still exist for this submission, table by table."""
    counts: dict[str, int] = {}
    for label, query in (
        ("system_scores", "select count(*) from system_scores where submission_id = %s"),
        ("agent_traces", "select count(*) from agent_traces where submission_id = %s"),
        ("feedback", "select count(*) from feedback where submission_id = %s"),
        ("jobs", "select count(*) from jobs where submission_id = %s"),
        (
            "evidence",
            "select count(*) from evidence e join system_scores s on s.id = e.system_score_id "
            "where s.submission_id = %s",
        ),
    ):
        counts[label] = conn.execute(query, (submission_id,)).fetchone()[0]
    return counts


def test_assignment_list_reports_counts_and_latest_run(seeded):
    run_id = _run_batch(seeded)
    listed = seeded["client"].get("/assignments").json()

    assert [row["id"] for row in listed] == [str(seeded["assignment"])]
    entry = listed[0]
    assert entry["title"] == "Tugas Risiko"
    assert entry["course"] == "MRTI"
    assert entry["submission_count"] == 1
    assert entry["latest_run"]["id"] == run_id
    assert entry["latest_run"]["status"] == "done"
    assert entry["latest_run"]["jobs"] == {"done": 1}


def test_create_assignment_uses_the_newest_rubric(seeded):
    client, csrf = seeded["client"], seeded["csrf"]
    created = client.post(
        "/assignments",
        json={"title": "Tugas Kedua", "course": "MRTI", "prompt_text": "Apa itu risiko?"},
        headers={"X-CSRF-Token": csrf},
    )
    assert created.status_code == 201, created.text
    body = created.json()
    assert body["submission_count"] == 0
    assert body["latest_run"] is None

    with db.tenant_tx(seeded["user"]) as conn:
        rubric = conn.execute(
            "select rv.version from assignments a join rubric_versions rv on rv.id = a.rubric_version_id "
            "where a.id = %s",
            (body["id"],),
        ).fetchone()[0]
    assert rubric == 1

    detail = client.get(f"/assignments/{body['id']}").json()
    assert detail["submissions"] == []
    assert detail["runs"] == []


def test_create_assignment_without_a_rubric_is_a_409(client):
    client.post("/auth/signup", json=OTHER_LECTURER)
    response = client.post(
        "/assignments", json={"title": "Tugas", "course": "MRTI"}, headers=_csrf(client)
    )
    assert response.status_code == 409
    assert response.json()["detail"] == (
        "No rubric version is seeded for this account; run scripts/seed_rubric.py first."
    )


def test_detail_shows_the_latest_run_and_its_rubric(seeded):
    first = _run_batch(seeded)
    second = _run_batch(seeded)

    response = seeded["client"].get(f"/submissions/{seeded['submission']}")
    assert response.status_code == 200, response.text
    body = response.json()

    assert body["run"] == {"id": second, "status": "done", "mode": "multi_agent"}
    assert body["run"]["id"] != first, "the detail view must follow the newest run, not the first"
    assert body["assignment_id"] == str(seeded["assignment"]), "the detail screen links back to it"
    assert body["rubric"] == {
        "version": "1",
        "levels": LEVELS,
        "indicators": INDICATORS,
        "level_names": LEVEL_NAMES,
    }
    assert body["student_analysis_text"]
    assert body["pseudonym"].startswith("S-")

    indicators = {s["indicator"] for s in body["scores"] if s["dimension"] == "metacognitive"}
    assert indicators == set(INDICATORS)
    cognitive = next(s for s in body["scores"] if s["dimension"] == "cognitive")
    assert cognitive["level"] == "C4"
    assert cognitive["rationale"]
    assert cognitive["evidence"][0]["quote"] == QUOTE_COGNITIVE
    assert body["feedback"]["generated_text"]
    assert body["feedback"]["edited_text"] is None
    assert body["reviews"] == []

    assignment = seeded["client"].get(f"/assignments/{seeded['assignment']}").json()
    assert [run["id"] for run in assignment["runs"]] == [second, first]
    assert assignment["latest_run"]["id"] == second
    assert assignment["submission_count"] == 1
    assert [s["id"] for s in assignment["submissions"]] == [str(seeded["submission"])]
    assert assignment["submissions"][0]["analysis_chars"] > 0
    assert assignment["submissions"][0]["review_flags"] == []


def test_detail_of_an_unscored_submission_is_not_an_error(seeded):
    body = seeded["client"].get(f"/submissions/{seeded['submission']}").json()
    assert body["run"] is None
    assert body["rubric"] is None
    assert body["scores"] == []
    assert body["feedback"] is None


def test_the_results_payload_previews_exactly_what_the_export_writes(seeded):
    """The screen previews the file, so both must read the same two things.

    The results payload used to send the *generated* feedback while the export wrote the lecturer's edit,
    and it carried no review note at all — a preview built on it would have promised something the file did
    not contain. This asserts the preview and the file agree, field by field.
    """
    client = seeded["client"]
    run_id = _run_batch(seeded)
    score_id = _score_id(seeded["user"], run_id, "cognitive")
    client.post(
        f"/system_scores/{score_id}/review",
        json={"action": "override", "new_level": "C5", "reason": "Bukti kuat."},
        headers=_csrf(client),
    )
    client.patch(
        f"/submissions/{seeded['submission']}/feedback",
        json={"edited_text": "Versi dosen."},
        headers=_csrf(client),
    )

    submission = client.get(f"/runs/{run_id}/results").json()["submissions"][0]
    cognitive = next(s for s in submission["scores"] if s["dimension"] == "cognitive")
    assert cognitive["review"] is not None
    assert (
        cognitive["review"]["action"],
        cognitive["review"]["new_level"],
        cognitive["review"]["reason"],
    ) == ("override", "C5", "Bukti kuat.")
    assert all(
        score["review"] is None
        for score in submission["scores"]
        if score["dimension"] == "metacognitive"
    ), "a score nobody reviewed carries no note"

    assert submission["feedback_text"] == "Versi dosen."
    assert submission["feedback_edited"] is True

    rows = list(
        csv.reader(io.StringIO(client.get(f"/runs/{run_id}/export.csv").content.decode("utf-8-sig")))
    )
    header = rows[0]
    row = next(r for r in rows[1:] if r[header.index("dimension")] == "cognitive")
    assert row[header.index("feedback_text")] == submission["feedback_text"]
    assert row[header.index("review_action")] == cognitive["review"]["action"]
    assert row[header.index("review_reason")] == cognitive["review"]["reason"]


def test_a_rubric_without_level_names_reports_none_rather_than_inventing_them(seeded):
    """A rubric row seeded before names existed carries none.

    The payload then reports an empty map and the screen shows the bare code: a name the data does not have
    is not the app's to make up (PRD §13 keeps the level scale in the rubric row).
    """
    client = seeded["client"]
    run_id = _run_batch(seeded)
    with db.tenant_tx(seeded["user"]) as conn:
        conn.execute("update rubric_versions set descriptors = '{}'::jsonb")

    assert client.get(f"/runs/{run_id}/results").json()["level_names"] == {}
    assert client.get(f"/submissions/{seeded['submission']}").json()["rubric"]["level_names"] == {}


def test_override_needs_a_reason_and_a_known_level(seeded):
    client, csrf = seeded["client"], seeded["csrf"]
    run_id = _run_batch(seeded)
    score_id = _score_id(seeded["user"], run_id)
    url = f"/system_scores/{score_id}/review"

    no_reason = client.post(
        url, json={"action": "override", "new_level": "C2"}, headers={"X-CSRF-Token": csrf}
    )
    assert no_reason.status_code == 422, no_reason.text
    blank_reason = client.post(
        url, json={"action": "override", "new_level": "C2", "reason": "   "}, headers={"X-CSRF-Token": csrf}
    )
    assert blank_reason.status_code == 422, blank_reason.text
    invented = client.post(
        url, json={"action": "override", "new_level": "A1", "reason": "Bukan level rubrik."},
        headers={"X-CSRF-Token": csrf},
    )
    assert invented.status_code == 422, invented.text

    accepted = client.post(url, json={"action": "accept"}, headers={"X-CSRF-Token": csrf})
    assert accepted.status_code == 201, accepted.text
    assert accepted.json()["new_level"] is None
    assert accepted.json()["reason"] is None

    overridden = client.post(
        url,
        json={"action": "override", "new_level": "C2", "reason": "Analisisnya baru memahami."},
        headers={"X-CSRF-Token": csrf},
    )
    assert overridden.status_code == 201, overridden.text
    assert overridden.json()["new_level"] == "C2"

    with db.tenant_tx(seeded["user"]) as conn:
        score = conn.execute(
            "select level, rationale from system_scores where id = %s", (score_id,)
        ).fetchone()
        reviews = conn.execute(
            "select action, new_level from lecturer_reviews where system_score_id = %s "
            "order by reviewed_at",
            (score_id,),
        ).fetchall()
    # FR-11: the system score is the system's, however the lecturer judged it.
    assert score[0] == "C4"
    assert score[1] == "Mahasiswa menilai dua kerangka kerja dan membandingkan fokusnya."
    assert [row[0] for row in reviews] == ["accept", "override"]

    detail = seeded["client"].get(f"/submissions/{seeded['submission']}").json()
    assert detail["reviews"] == [
        {
            "system_score_id": str(score_id),
            "action": "override",
            "new_level": "C2",
            "reason": "Analisisnya baru memahami.",
            "reviewed_at": overridden.json()["reviewed_at"],
        }
    ]


def test_feedback_edit_and_blank_clear(seeded):
    client, csrf = seeded["client"], seeded["csrf"]
    url = f"/submissions/{seeded['submission']}/feedback"

    early = client.patch(url, json={"edited_text": "Belum ada."}, headers={"X-CSRF-Token": csrf})
    assert early.status_code == 404, early.text

    _run_batch(seeded)
    edited = client.patch(
        url, json={"edited_text": "  Umpan balik yang saya rapikan.  "}, headers={"X-CSRF-Token": csrf}
    )
    assert edited.status_code == 200, edited.text
    assert edited.json()["edited_text"] == "Umpan balik yang saya rapikan."
    assert edited.json()["generated_text"], "the generated text is kept beside the edit"
    assert edited.json()["edited_by"] == str(seeded["user"])

    cleared = client.patch(url, json={"edited_text": "   "}, headers={"X-CSRF-Token": csrf})
    assert cleared.status_code == 200, cleared.text
    assert cleared.json()["edited_text"] is None


# The research dataset's own column order. Spelled out rather than imported: a test that reads the
# header from the code it checks would pass after any reordering (FR-16).
EXPECTED_HEADER = [
    "assignment_title",
    "submission_pseudonym",
    "dimension",
    "indicator",
    "level",
    "rationale",
    "evidence_quote",
    "start_offset",
    "end_offset",
    "review_action",
    "review_new_level",
    "review_reason",
    "reviewed_at",
    "feedback_text",
    "run_id",
    "mode",
    "model",
    "temperature",
    "seed",
    "rubric_version",
    "started_at",
    "finished_at",
]


def test_export_csv_and_xlsx(seeded):
    client, csrf = seeded["client"], seeded["csrf"]
    run_id = _run_batch(seeded)
    score_id = _score_id(seeded["user"], run_id, "metacognitive")
    client.post(
        f"/system_scores/{score_id}/review",
        json={"action": "override", "new_level": "C5", "reason": "Sudah terstruktur."},
        headers={"X-CSRF-Token": csrf},
    )
    client.patch(
        f"/submissions/{seeded['submission']}/feedback",
        json={"edited_text": "Versi dosen."},
        headers={"X-CSRF-Token": csrf},
    )

    response = client.get(f"/runs/{run_id}/export.csv")
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("text/csv")
    assert response.headers["content-disposition"] == f'attachment; filename="run-{run_id}-scores.csv"'

    text = response.content.decode("utf-8")
    assert text.startswith("\ufeff"), "Excel needs the BOM to read Indonesian text"
    rows = list(csv.reader(io.StringIO(text.removeprefix("\ufeff"))))
    assert rows[0] == EXPECTED_HEADER
    assert COLUMNS == EXPECTED_HEADER

    # One row per evidence span, plus one row for a score whose quote failed the evidence check.
    with db.tenant_tx(seeded["user"]) as conn:
        expected_rows = conn.execute(
            "select count(*) from system_scores sc left join evidence e on e.system_score_id = sc.id "
            "where sc.run_id = %s",
            (run_id,),
        ).fetchone()[0]
        spans = conn.execute(
            "select count(*) from evidence e join system_scores sc on sc.id = e.system_score_id "
            "where sc.run_id = %s",
            (run_id,),
        ).fetchone()[0]
    assert spans > 0
    assert len(rows) - 1 == expected_rows > spans, "a score without a span still exports a row"

    by_pseudonym = {row[1] for row in rows[1:]}
    assert len(by_pseudonym) == 1
    cognitive_row = next(row for row in rows[1:] if row[2] == "cognitive")
    assert cognitive_row[0] == "Tugas Risiko"
    assert cognitive_row[6] == QUOTE_COGNITIVE
    assert cognitive_row[13] == "Versi dosen."
    assert cognitive_row[14] == run_id
    assert cognitive_row[15] == "multi_agent"
    assert cognitive_row[19] == "1"
    # The metadata columns are the *run's* snapshot, not the current environment: read them back from
    # the run row, so the assertion still says something true where no provider is configured (CI) and
    # actually checks the stronger claim — that the export reports what the run froze, not what the
    # process happens to be configured with now.
    with db.tenant_tx(seeded["user"]) as conn:
        config = conn.execute("select model_config from runs where id = %s", (run_id,)).fetchone()[0]
    assert cognitive_row[16] == (config.get("model") or "")
    assert cognitive_row[17] == ("" if config.get("temperature") is None else str(config["temperature"]))
    assert cognitive_row[18] == ("" if config.get("seed") is None else str(config["seed"]))

    reviewed = [
        row for row in rows[1:] if row[9] not in ("", None)
    ]
    assert [(row[9], row[10], row[11]) for row in reviewed] == [
        ("override", "C5", "Sudah terstruktur.")
    ]

    # A score with no located quote keeps its row, with empty offsets.
    empty = next(row for row in rows[1:] if row[6] == "" and row[2] == "metacognitive")
    assert empty[7] == "" and empty[8] == ""

    book = openpyxl.load_workbook(io.BytesIO(client.get(f"/runs/{run_id}/export.xlsx").content))
    assert book.sheetnames == ["Scores"]
    assert [cell.value for cell in book["Scores"][1]] == EXPECTED_HEADER
    assert book["Scores"].max_row - 1 == expected_rows


def test_deleting_a_submission_takes_everything_derived_from_it(seeded):
    """One delete, because the rows a lecturer can no longer see are still rows about a student."""
    client, csrf = seeded["client"], seeded["csrf"]
    run_id = _run_batch(seeded)
    submission = uuid.UUID(seeded["submission"])
    with db.tenant_tx(seeded["user"]) as conn:
        counts = _rows_for(conn, submission)
    assert counts["system_scores"] > 0 and counts["agent_traces"] > 0 and counts["jobs"] > 0
    # The filename ("Tugas 1_Dosen Runs.docx") carries a name, so this submission has one — not a
    # foreign key any more (DECISIONS.md, 2026-10-16), so it needs its own before/after check.
    assert identity_store.read(submission, seeded["user"]) == "Dosen Runs"

    response = client.delete(f"/submissions/{submission}", headers={"X-CSRF-Token": csrf})
    assert response.status_code == 204, response.text

    assert client.get(f"/submissions/{submission}").status_code == 404
    with db.tenant_tx(seeded["user"]) as conn:
        assert _rows_for(conn, submission) == {
            "system_scores": 0,
            "evidence": 0,
            "agent_traces": 0,
            "feedback": 0,
            "jobs": 0,
        }
    assert identity_store.read(submission, seeded["user"]) is None
    assignment = client.get(f"/assignments/{seeded['assignment']}").json()
    assert assignment["submissions"] == []
    assert client.get(f"/runs/{run_id}/results").json()["submissions"] == []


def test_a_submission_being_graded_cannot_be_deleted_under_the_worker(seeded):
    """The worker holds that submission's text: deleting the row mid-flight would make it fail to store a
    result and take the whole pass down with it."""
    client, csrf = seeded["client"], seeded["csrf"]
    run_id = client.post(
        f"/assignments/{seeded['assignment']}/runs", headers={"X-CSRF-Token": csrf}
    ).json()["id"]
    with db.tenant_tx(seeded["user"]) as conn:
        assert len(claim_jobs(conn, owner_id=seeded["user"], run_id=run_id, limit=10)) == 1

    response = client.delete(f"/submissions/{seeded['submission']}", headers={"X-CSRF-Token": csrf})
    assert response.status_code == 409, response.text
    assert client.get(f"/submissions/{seeded['submission']}").status_code == 200


def test_the_research_export_carries_both_modes_with_the_metadata_to_tell_them_apart(seeded):
    """FR-16: the thesis joins this file with the two-lecturer rating sheet.

    A per-run file would force a join of files before the join with the sheet, so this one covers the whole
    assignment with `mode` and the run's own snapshot as columns.
    """
    client = seeded["client"]
    multi = _run_batch(seeded)
    baseline = _run_batch(seeded, mode="single_agent")
    score_id = _score_id(seeded["user"], multi, "cognitive")
    client.post(
        f"/system_scores/{score_id}/review",
        json={"action": "override", "new_level": "C5", "reason": "Bukti pembanding."},
        headers=_csrf(client),
    )

    response = client.get(f"/assignments/{seeded['assignment']}/research-export.csv")
    assert response.status_code == 200, response.text
    rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
    assert rows[0] == RESEARCH_COLUMNS
    header = rows[0]
    by_mode = {
        row[header.index("mode")]: row
        for row in rows[1:]
        if row[header.index("dimension")] == "cognitive"
    }
    assert set(by_mode) == {"multi_agent", "single_agent"}, by_mode
    assert by_mode["multi_agent"][header.index("run_id")] == multi
    assert by_mode["single_agent"][header.index("run_id")] == baseline
    assert by_mode["multi_agent"][header.index("review_action")] == "override"
    assert by_mode["multi_agent"][header.index("review_reason")] == "Bukti pembanding."
    assert by_mode["single_agent"][header.index("review_action")] == ""

    # `feedback` is scoped to the submission, not the run, so this column reads the run's own synthesis
    # trace rather than that mutable table — otherwise a later run's feedback (or its absence) would show
    # up against an earlier run's row.
    assert by_mode["multi_agent"][header.index("feedback_text")] == (
        "Perbandingan Anda sudah baik; tambahkan contoh penerapan agar lebih kuat."
    )
    assert by_mode["single_agent"][header.index("feedback_text")] == "", "FR-15: baseline has no feedback"

    # The metadata columns come from each run's own snapshot, so the file explains what produced a score.
    with db.tenant_tx(seeded["user"]) as conn:
        config = conn.execute(
            "select model_config from runs where id = %s", (multi,)
        ).fetchone()[0]
    assert by_mode["multi_agent"][header.index("model")] == (config.get("model") or "")
    assert by_mode["multi_agent"][header.index("rubric_version")] == "1"
    assert by_mode["multi_agent"][header.index("run_status")] == "done"

    # The metacognitive indicators are rows too, one per indicator per mode.
    indicators = [
        row[header.index("indicator")] for row in rows[1:] if row[header.index("mode")] == "single_agent"
    ]
    assert sorted(indicators) == sorted([*INDICATORS, ""]), indicators

    book = openpyxl.load_workbook(
        io.BytesIO(client.get(f"/assignments/{seeded['assignment']}/research-export.xlsx").content)
    )
    assert [cell.value for cell in book["Scores"][1]] == RESEARCH_COLUMNS
    assert book["Scores"].max_row - 1 == len(rows) - 1


def test_a_baseline_run_does_not_erase_an_earlier_multi_agent_runs_feedback(seeded):
    """`feedback` is keyed by submission, not by run (no column for it); a baseline run afterwards on the
    same submission must leave the multi-agent run's own feedback exactly as it was, not clear it for a
    mode that never writes it.
    """
    client = seeded["client"]
    multi = _run_batch(seeded)
    before = client.get(f"/runs/{multi}/results").json()["submissions"][0]["feedback_text"]
    assert before, "the seeded multi-agent run must produce feedback for this test to mean anything"

    _run_batch(seeded, mode="single_agent")

    after = client.get(f"/runs/{multi}/results").json()["submissions"][0]["feedback_text"]
    assert after == before


def test_the_export_is_only_this_lecturers(seeded):
    """Both export families refuse another lecturer's assignment or run."""
    other = seeded["client"]
    other.post("/auth/signup", json=OTHER_LECTURER)
    assert other.get(f"/assignments/{seeded['assignment']}/research-export.csv").status_code == 404
    assert other.get(f"/assignments/{seeded['assignment']}/research-export.xlsx").status_code == 404


def test_another_lecturer_reaches_nothing(seeded):
    client = seeded["client"]
    run_id = _run_batch(seeded)
    score_id = _score_id(seeded["user"], run_id)

    client.post("/auth/signup", json=OTHER_LECTURER)
    headers = _csrf(client)

    assert client.get("/assignments").json() == []
    assert client.get(f"/assignments/{seeded['assignment']}").status_code == 404
    assert client.get(f"/submissions/{seeded['submission']}").status_code == 404
    assert client.get(f"/runs/{run_id}").status_code == 404
    assert client.get(f"/runs/{run_id}/results").status_code == 404
    assert client.get(f"/runs/{run_id}/export.csv").status_code == 404
    assert client.get(f"/runs/{run_id}/export.xlsx").status_code == 404
    assert (
        client.post(
            f"/system_scores/{score_id}/review", json={"action": "accept"}, headers=headers
        ).status_code
        == 404
    )
    assert (
        client.patch(
            f"/submissions/{seeded['submission']}/feedback",
            json={"edited_text": "Punya saya."},
            headers=headers,
        ).status_code
        == 404
    )
    assert (
        client.post(
            f"/assignments/{seeded['assignment']}/runs", headers=headers
        ).status_code
        == 404
    )

    with db.tenant_tx(seeded["user"]) as conn:
        assert conn.execute("select count(*) from lecturer_reviews").fetchone()[0] == 0
        assert conn.execute("select edited_text from feedback").fetchone()[0] is None
