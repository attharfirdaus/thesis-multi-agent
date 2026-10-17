"""Phase-3 gate at the API level: a synthetic submission imports, its identifiers are gone,
the boundary can be corrected, and one bad file does not fail the batch.

The fixture document deliberately reproduces the real files' shape — a two-column table repeated per
prompt, with the student's name and NIM inside the text — so the assertions are about the actual
hazard rather than a convenient one.
"""

import io
import json
import uuid

import docx
import pytest
from psycopg.rows import dict_row

from app import db
from app.ingest import identity_store
from app.ingest.segment import NO_AI_WARNING
from app.llm.client import Completion

NAME = "Budi Santoso"
NIM = "28123456"


@pytest.fixture
def lecturer(client):
    client.post(
        "/auth/signup",
        json={"email": "ingest@example.ac.id", "name": "Dosen Ing", "password": "kata-sandi-panjang-2026"},
    )
    user = uuid.UUID(client.get("/auth/me").json()["user"]["id"])
    with db.tenant_tx(user) as conn:
        assignment = conn.execute(
            "insert into assignments (owner_id, title, course) values (%s, %s, %s) returning id",
            (user, "Tugas Risiko", "MRTI"),
        ).fetchone()[0]
    return {"client": client, "user": user, "assignment": assignment, "csrf": client.cookies.get("ta_csrf")}


def build_submission_docx() -> bytes:
    """The real shape: | Prompt | Hasil | then | Critical Review | <analysis>."""
    document = docx.Document()
    table = document.add_table(rows=0, cols=2)
    for label, content in [
        ("Prompt", "Apa yang dimaksud COSO pada manajemen risiko?"),
        ("Hasil", "COSO adalah kerangka kerja yang dikembangkan oleh Committee of Sponsoring Organization."),
        (
            "Critical Review",
            f"Dari hasil prompt di atas, saya memahami bahwa COSO adalah framework pengendalian internal. "
            f"Nama : {NAME}\nNIM : {NIM}\n"
            f"Perbedaan mendasar dengan ISO adalah COSO berfokus pada kontrol internal. "
            f"Hubungi saya di budi.santoso@example.ac.id bila perlu.",
        ),
    ]:
        row = table.add_row().cells
        row[0].text = label
        row[1].text = content
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def upload(lecturer, *files) -> dict:
    return lecturer["client"].post(
        f"/assignments/{lecturer['assignment']}/submissions",
        files=list(files),
        headers={"X-CSRF-Token": lecturer["csrf"]},
    )


DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def test_import_removes_identifiers_and_keeps_the_mapping(lecturer):
    response = upload(lecturer, ("files", ("Tugas 1_Budi Santoso.docx", build_submission_docx(), DOCX_MIME)))
    assert response.status_code == 201, response.text
    entry = response.json()["files"][0]

    assert entry["status"] == "imported", entry.get("error") or entry
    assert entry["segmentation_method"] == "label"
    assert entry["has_ai_output"] is True
    assert entry["removed_by_class"].get("name_label", 0) >= 1
    assert entry["removed_by_class"].get("student_id", 0) >= 1

    # Read back through the same tenant context the app uses: the tables have FORCE ROW LEVEL
    # SECURITY, so a connection without `app.current_user_id` sees nothing at all.
    with db.tenant_tx(lecturer["user"]) as conn, conn.cursor(row_factory=dict_row) as cur:
        stored = cur.execute(
            "select ai_output_text, student_analysis_text, segmentation_status "
            "from submissions where id = %s",
            (entry["submission_id"],),
        ).fetchone()
        identity = identity_store.read(uuid.UUID(entry["submission_id"]), lecturer["user"])
        events = cur.execute(
            "select cls, count from redaction_events where submission_id = %s",
            (entry["submission_id"],),
        ).fetchall()

    # The AI answer and the student's own analysis must have landed in different fields.
    assert "Committee of Sponsoring Organization" in stored["ai_output_text"]
    assert "saya memahami bahwa COSO" in stored["student_analysis_text"]
    assert "Committee of Sponsoring Organization" not in stored["student_analysis_text"]

    # The point of FR-04: the identifiers are gone from the stored text...
    assert NAME not in stored["student_analysis_text"]
    assert NAME not in stored["ai_output_text"]
    assert NIM not in stored["student_analysis_text"]
    assert "budi.santoso@example.ac.id" not in stored["student_analysis_text"]
    # ...while the lecturer can still tell whose work it is.
    assert identity == NAME
    assert {e["cls"] for e in events} >= {"name_label", "student_id", "email"}
    # No row anywhere may hold the removed value.
    assert all(NAME not in str(dict(e)) for e in events)


def test_a_document_without_ai_output_records_that_once(lecturer):
    """One fact, one warning.

    The router used to append its own phrasing of the fact the segmenter had already reported, so the
    lecturer's screen showed the same sentence twice (seen live, 2026-10-16).
    """
    document = docx.Document()
    document.add_paragraph("Analisis saya tentang manajemen risiko, tanpa keluaran AI sama sekali.")
    buffer = io.BytesIO()
    document.save(buffer)

    entry = upload(
        lecturer, ("files", ("Tugas tanpa AI.docx", buffer.getvalue(), DOCX_MIME))
    ).json()["files"][0]

    assert entry["has_ai_output"] is False
    assert [warning for warning in entry["warnings"] if "AI" in warning] == [NO_AI_WARNING], entry


def test_one_bad_file_does_not_fail_the_batch(lecturer):
    response = upload(
        lecturer,
        ("files", ("catatan.md", b"# not supported", "text/markdown")),
        ("files", ("Tugas 2_Budi Santoso.docx", build_submission_docx(), DOCX_MIME)),
    )
    assert response.status_code == 201, response.text
    files = {f["filename"]: f for f in response.json()["files"]}
    assert files["catatan.md"]["status"] == "failed"
    assert "md" in (files["catatan.md"]["error"] or "")
    assert files["Tugas 2_Budi Santoso.docx"]["status"] == "imported", files["Tugas 2_Budi Santoso.docx"]


def test_listing_shows_the_submission(lecturer):
    upload(lecturer, ("files", ("Tugas 3_Budi Santoso.docx", build_submission_docx(), DOCX_MIME)))
    listed = lecturer["client"].get(f"/assignments/{lecturer['assignment']}/submissions")
    assert listed.status_code == 200
    rows = listed.json()
    assert len(rows) == 1
    assert rows[0]["has_ai_output"] is True
    assert rows[0]["analysis_chars"] > 0
    assert rows[0]["pseudonym"].startswith("S-")


def test_lecturer_can_correct_the_boundary(lecturer):
    entry = upload(
        lecturer, ("files", ("Tugas 4_Budi Santoso.docx", build_submission_docx(), DOCX_MIME))
    ).json()["files"][0]

    # "None of this was AI output" — the lecturer moves the boundary to the very start.
    patched = lecturer["client"].patch(
        f"/submissions/{entry['submission_id']}/segmentation",
        json={"analysis_start": 0},
        headers={"X-CSRF-Token": lecturer["csrf"]},
    )
    assert patched.status_code == 200, patched.text
    body = patched.json()
    assert body["segmentation_status"] == "confirmed"
    assert body["has_ai_output"] is False

    # And a boundary past the end is refused rather than trimming silently.
    too_far = lecturer["client"].patch(
        f"/submissions/{entry['submission_id']}/segmentation",
        json={"analysis_start": 10_000_000},
        headers={"X-CSRF-Token": lecturer["csrf"]},
    )
    assert too_far.status_code == 422


def test_another_lecturer_cannot_see_or_touch_the_submission(lecturer):
    entry = upload(
        lecturer, ("files", ("Tugas 5_Budi Santoso.docx", build_submission_docx(), DOCX_MIME))
    ).json()["files"][0]

    other = lecturer["client"]
    other.post("/auth/logout", headers={"X-CSRF-Token": other.cookies.get("ta_csrf")})
    other.post(
        "/auth/signup",
        json={"email": "lain@example.ac.id", "name": "Dosen Lain", "password": "kata-sandi-panjang-2026"},
    )

    assert other.get(f"/assignments/{lecturer['assignment']}/submissions").json() == []
    assert other.patch(
        f"/submissions/{entry['submission_id']}/segmentation",
        json={"analysis_start": 1},
        headers={"X-CSRF-Token": other.cookies.get("ta_csrf")},
    ).status_code == 404


def test_an_unlabelled_document_falls_through_to_the_llm_and_the_result_is_stored(lecturer, monkeypatch):
    """FR-05, end to end through the real upload endpoint: no table, no recognised label, so the rule-
    based segmenter reports `method="none"` and the LLM fallback is the only way this gets a confident
    split. The caller must still never send an unredacted identifier — this document carries one.
    """
    ai_text = "Berikut adalah definisi risiko menurut ISO 31000 beserta komponennya."
    analysis_text = f"Analisis saya terhadap definisi ini memerlukan konteks organisasi.\nNama : {NAME}"
    document = docx.Document()
    document.add_paragraph(ai_text)
    document.add_paragraph(analysis_text)
    buffer = io.BytesIO()
    document.save(buffer)

    class StubProviderClient:
        def __init__(self) -> None:
            self.calls = 0

        def complete_json(self, *, system, user, schema, **_kwargs):
            self.calls += 1
            assert NAME not in system and NAME not in user, "an unredacted identifier reached the call"
            return Completion(
                parsed=schema.model_validate(
                    json.loads(json.dumps({"has_ai_output": True, "ai_output_quote": ai_text}))
                ),
                raw="{}",
                latency_ms=5,
                attempts=1,
            )

    stub = StubProviderClient()
    monkeypatch.setattr("app.ingest.llm_segment.ProviderClient", lambda: stub)

    entry = upload(
        lecturer,
        ("files", ("tanpa_label.docx", buffer.getvalue(), DOCX_MIME)),
    ).json()["files"][0]

    assert entry["status"] == "imported", entry.get("error") or entry
    assert entry["segmentation_method"] == "llm"
    assert stub.calls == 1

    with db.tenant_tx(lecturer["user"]) as conn, conn.cursor(row_factory=dict_row) as cur:
        stored = cur.execute(
            "select ai_output_text, student_analysis_text from submissions where id = %s",
            (entry["submission_id"],),
        ).fetchone()
    assert stored["ai_output_text"] == ai_text
    assert stored["student_analysis_text"] == (
        "Analisis saya terhadap definisi ini memerlukan konteks organisasi.\nNama : [REDACTED]"
    )
    assert NAME not in stored["ai_output_text"]
    assert NAME not in stored["student_analysis_text"]
